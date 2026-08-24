"""Convert PROJECT_REPORT.md into a submission-ready Word document.

Usage:  python generate_docx.py
Output: Project_Report.docx
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

SRC = Path(__file__).parent / "PROJECT_REPORT.md"
OUT = Path(__file__).parent / "Project_Report.docx"

BRAND = RGBColor(0x1E, 0x2A, 0x4A)
ACCENT = RGBColor(0x3B, 0x5B, 0xDB)
MUTED = RGBColor(0x5B, 0x64, 0x78)
CODE_BG = "F2F4FA"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_runs(paragraph, text, size=None, color=None, italic=False, bold=False):
    """Parse **bold**, *italic* and `code` inline markers into styled runs."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt((size or 11) - 0.5)
            run.font.color.rgb = ACCENT
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        if italic:
            run.italic = True
        if bold:
            run.bold = True


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell_text in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        add_runs(p, cell_text.strip(), size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shd = cell._tc.get_or_add_tcPr().makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd", {})
        shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val", "clear")
        shd.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill", "1E2A4A")
        cell._tc.get_or_add_tcPr().append(shd)
    for i, row in enumerate(body, start=1):
        for j, cell_text in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, cell_text.strip(), size=9.5)
    doc.add_paragraph()


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    n = len(lines)
    title_done = False

    while i < n:
        line = lines[i].rstrip()

        # fenced code block
        if line.startswith("```"):
            buf = []
            i += 1
            while i < n and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x14, 0x1A, 0x2E)
            continue

        # horizontal rule
        if re.match(r"^\s*---+\s*$", line):
            i += 1
            continue

        # markdown table
        if line.lstrip().startswith("|") and i + 1 < n and re.match(
                r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            rows = []
            while i < n and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2).strip()
            if level == 1 and not title_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(24)
                p.paragraph_format.space_after = Pt(4)
                add_runs(p, text, size=22, bold=True, color=BRAND)
                title_done = True
            elif level == 1:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(20)
                p.paragraph_format.space_after = Pt(8)
                add_runs(p, text, size=17, bold=True, color=BRAND)
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                add_runs(p, text, size=14, bold=True, color=BRAND)
            elif level == 3:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                add_runs(p, text, size=12, bold=True, color=ACCENT)
            else:
                p = doc.add_paragraph()
                add_runs(p, text, size=11, bold=True, italic=True)
            i += 1
            continue

        # blockquote
        if line.lstrip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(8)
            add_runs(p, line.lstrip()[1:].strip(), size=11, italic=True,
                     color=MUTED)
            i += 1
            continue

        # bullet list (including nested)
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent, text = len(m.group(1)), m.group(2).strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * (indent // 2))
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, "•  " + text, size=11)
            i += 1
            continue

        # numbered list (supports "1." and nested "1.1" prefixes)
        m = re.match(r"^(\s*)([\d.]+)\s+(.*)$", line)
        if m and re.match(r"^\d+(\.\d+)*\.?$", m.group(2)):
            indent, num, text = len(m.group(1)), m.group(2).strip(), m.group(3).strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3 + 0.25 * (indent // 2))
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, f"{num} {text}", size=11)
            i += 1
            continue

        # empty line -> spacing only
        if not line.strip():
            i += 1
            continue

        # plain paragraph (allow multi-line continuation)
        buf = [line]
        i += 1
        while i < n and lines[i].strip() and not lines[i].lstrip().startswith(("#", "|", "-", "*", ">", "```")):
            if re.match(r"^\s*\d+(\.\d+)*\.?\s+", lines[i]):
                break
            buf.append(lines[i].rstrip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        add_runs(p, " ".join(x.strip() for x in buf if x.strip()), size=11)

    doc.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
